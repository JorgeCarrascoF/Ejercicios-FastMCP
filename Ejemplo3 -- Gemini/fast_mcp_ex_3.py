from fastmcp import FastMCP
from fpdf import FPDF
import os
import pypdf
from docx import Document
import sys

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
server = FastMCP("TextSintetizer")

def get_safe_path(filename_or_path: str) -> str:
    clean_filename = os.path.basename(filename_or_path)
    return os.path.join(BASE_DIR, clean_filename)

@server.tool()
def read_local_file(filepath: str) -> str:
    safe_path = get_safe_path(filepath)
    sys.stderr.write(f"[DEBUG] Leyendo archivo: {safe_path}\n")

    if not os.path.exists(safe_path):
        return f"Error: No encuentro el archivo '{os.path.basename(safe_path)}' en /app."
    
    _, ext = os.path.splitext(safe_path)
    ext = ext.lower()

    try:
        if ext == '.pdf':
            reader = pypdf.PdfReader(safe_path)
            text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text.append(extracted)
            return "\n".join(text)
        elif ext == '.docx':
            doc = Document(safe_path)
            return "\n".join([para.text for para in doc.paragraphs])
        else:
            with open(safe_path, 'r', encoding='utf-8') as f:
                return f.read()
    except Exception as e:
        return f"Error leyendo: {str(e)}"

@server.tool()
def create_pdf_summary(summary_text: str, filename: str = "resumen.pdf") -> str:
    safe_path = get_safe_path(filename)
    sys.stderr.write(f"[DEBUG] Generando PDF en: {safe_path}\n")

    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, "Resumen IA", ln=True, align='C')
        pdf.ln(10)
        pdf.set_font("Arial", size=12)
        
        replacements = {'\u2018': "'", '\u2019': "'", '\u201c': '"', '\u201d': '"', '—': '-'}
        for k, v in replacements.items():
            summary_text = summary_text.replace(k, v)
            
        clean_text = summary_text.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 7, clean_text)
        pdf.output(safe_path)
        
        return f"ÉXITO: PDF guardado como '{os.path.basename(safe_path)}'"
    except Exception as e:
        return f"Error: {str(e)}"

if __name__ == "__main__":
    server.run()